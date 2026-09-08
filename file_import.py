"""
多格式题目导入模块
支持：Excel(.xlsx)、Word(.docx)、PDF(.pdf)、图片(.png/.jpg)
"""
import os
import re
import json

from pdf_import import (
    parse_questions_from_text, parse_questions_with_numbers,
    parse_answers_from_text, import_dual_pdf, extract_text_from_pdf,
    normalize_answer, clean_analysis
)


# ============ Excel 导入 ============

EXCEL_HEADERS = ["题型", "题干", "选项A", "选项B", "选项C", "选项D", "选项E", "选项F", "答案", "解析", "科目"]

TYPE_MAP_CN = {
    "单选": "single", "单选题": "single", "single": "single",
    "多选": "multiple", "多选题": "multiple", "multiple": "multiple",
    "判断": "judge", "判断题": "judge", "judge": "judge",
    "案例": "case", "案例分析": "case", "案例分析题": "case", "简答": "case", "case": "case",
}


def generate_excel_template(filepath):
    """生成 Excel 导入标准模板"""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    except ImportError:
        return False, "未安装 openpyxl，请先运行 pip install openpyxl"

    wb = Workbook()
    ws = wb.active
    ws.title = "题目导入模板"

    # 表头样式
    header_font = Font(bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    thin_border = Border(
        left=Side(style='thin'), right=Side(style='thin'),
        top=Side(style='thin'), bottom=Side(style='thin')
    )

    # 写表头
    for col, header in enumerate(EXCEL_HEADERS, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
        cell.border = thin_border

    # 示例数据
    examples = [
        ["单选题", "Python中定义函数的关键字是？", "function", "def", "func", "define", "", "", "B", "Python使用def关键字定义函数", "Python基础"],
        ["多选题", "以下哪些是Python内置数据结构？", "list", "tuple", "dict", "set", "", "", "ABCD", "四种都是Python内置数据结构", "Python基础"],
        ["判断题", "Python字符串是可变对象。", "", "", "", "", "", "", "错误", "字符串是不可变对象", "Python基础"],
        ["案例分析题", "请简述Python中列表和元组的区别。", "", "", "", "", "", "", "列表是可变的，元组是不可变的...", "考察对核心数据结构的理解", "Python基础"],
    ]
    for row_idx, row_data in enumerate(examples, 2):
        for col_idx, value in enumerate(row_data, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.border = thin_border
            cell.alignment = Alignment(vertical="center", wrap_text=True)

    # 列宽
    col_widths = [10, 40, 15, 15, 15, 15, 12, 12, 10, 30, 12]
    for i, width in enumerate(col_widths, 1):
        ws.column_dimensions[chr(64 + i)].width = width

    # 第二个sheet：填写说明
    ws2 = wb.create_sheet("填写说明")
    instructions = [
        ["Excel 题目导入模板 - 填写说明"],
        [""],
        ["1. 题型：填写 单选题/多选题/判断题/案例分析题"],
        ["2. 题干：题目内容，必填"],
        ["3. 选项A~F：选择题的选项，判断题和案例题可留空"],
        ["4. 答案：单选填字母(如B)，多选填字母(如ABCD)，判断填 正确/错误，案例题填参考答案"],
        ["5. 解析：题目解析，可选"],
        ["6. 科目：题目所属科目，可选，不填则归入默认科目"],
        [""],
        ["注意：第一行是表头，不要修改或删除；从第二行开始填写题目。"],
    ]
    for row_idx, row_data in enumerate(instructions, 1):
        cell = ws2.cell(row=row_idx, column=1, value=row_data[0])
        if row_idx == 1:
            cell.font = Font(bold=True, size=14)
    ws2.column_dimensions['A'].width = 80

    wb.save(filepath)
    return True, "模板生成成功"


def import_excel(filepath):
    """从 Excel 导入题目"""
    try:
        from openpyxl import load_workbook
    except ImportError:
        return [], "未安装 openpyxl，请先运行 pip install openpyxl"

    wb = load_workbook(filepath, read_only=True, data_only=True)
    ws = wb.active

    questions = []
    rows = list(ws.iter_rows(min_row=2, values_only=True))

    for row in rows:
        if not row or all(c is None or str(c).strip() == "" for c in row):
            continue

        # 确保有足够的列
        row = list(row) + [None] * (len(EXCEL_HEADERS) - len(row))

        q_type_raw = str(row[0]).strip() if row[0] else "single"
        question_text = str(row[1]).strip() if row[1] else ""
        options = []
        for i in range(2, 8):  # 选项A-F
            if row[i] and str(row[i]).strip():
                options.append(str(row[i]).strip())
        answer = str(row[8]).strip() if row[8] else ""
        analysis = str(row[9]).strip() if row[9] else ""
        subject = str(row[10]).strip() if row[10] else ""

        if not question_text:
            continue

        q_type = TYPE_MAP_CN.get(q_type_raw, "single")

        # 判断题标准化
        if q_type == "judge":
            options = ["正确", "错误"]
            if answer in ['对', '正确', '√', 'T', 'True']:
                answer = '正确'
            elif answer in ['错', '错误', '×', 'F', 'False']:
                answer = '错误'

        # 多选题答案标准化
        if q_type == "multiple":
            letters = sorted(set(re.findall(r'[A-Z]', answer.upper())))
            answer = ''.join(letters) if letters else answer

        # 单选题答案标准化
        if q_type == "single":
            m = re.search(r'[A-Z]', answer.upper())
            answer = m.group(0) if m else answer

        questions.append({
            'q_type': q_type,
            'question_text': question_text,
            'options': options,
            'answer': answer,
            'analysis': analysis,
            'subject_name': subject,
            'source': 'Excel导入',
        })

    wb.close()
    return questions, f"成功解析 {len(questions)} 道题"


# ============ Word 导入 ============

def import_word(filepath):
    """从 Word 文档导入题目"""
    try:
        from docx import Document
    except ImportError:
        return [], "未安装 python-docx，请先运行 pip install python-docx"

    doc = Document(filepath)
    full_text = "\n".join([para.text for para in doc.paragraphs])

    # 也提取表格中的内容
    for table in doc.tables:
        for row in table.rows:
            row_text = " ".join([cell.text for cell in row.cells])
            full_text += "\n" + row_text

    questions = parse_questions_from_text(full_text)
    for q in questions:
        q['source'] = 'Word导入'

    return questions, f"成功解析 {len(questions)} 道题"


# ============ 图片导入 (OCR) ============

def import_image(filepath):
    """从图片导入题目（使用OCR识别）"""
    try:
        import pytesseract
        from PIL import Image
    except ImportError:
        return [], "未安装 OCR 依赖，请先运行 pip install pytesseract pillow，并安装 Tesseract OCR 引擎"

    try:
        img = Image.open(filepath)
        # 尝试中文识别，失败则用英文
        try:
            text = pytesseract.image_to_string(img, lang='chi_sim+eng')
        except Exception:
            text = pytesseract.image_to_string(img)

        questions = parse_questions_from_text(text)
        for q in questions:
            q['source'] = '图片OCR导入'

        return questions, f"OCR识别并解析 {len(questions)} 道题"
    except Exception as e:
        return [], f"图片识别失败：{str(e)}"


# ============ 统一导入入口 ============

def import_file(filepath, subject_id=0):
    """
    统一文件导入入口，根据扩展名自动选择解析方式
    返回 (题目列表, 消息)
    """
    ext = os.path.splitext(filepath)[1].lower()

    if ext == '.xlsx' or ext == '.xls':
        questions, msg = import_excel(filepath)
    elif ext == '.docx' or ext == '.doc':
        questions, msg = import_word(filepath)
    elif ext == '.pdf':
        from pdf_import import import_pdf
        questions, raw_text = import_pdf(filepath)
        msg = f"成功解析 {len(questions)} 道题"
    elif ext in ('.png', '.jpg', '.jpeg', '.bmp', '.tiff'):
        questions, msg = import_image(filepath)
    elif ext == '.json':
        with open(filepath, 'r', encoding='utf-8') as f:
            questions = json.load(f)
        msg = f"成功加载 {len(questions)} 道题"
    else:
        return [], f"不支持的文件格式：{ext}"

    # 分配科目
    for q in questions:
        if 'subject_id' not in q:
            q['subject_id'] = subject_id

    return questions, msg
