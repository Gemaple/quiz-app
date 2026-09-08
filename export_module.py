"""
导出模块：支持 Word / PDF / Excel 导出
V9 新增功能：
1. 考试记录导出（Word/PDF）：试卷形式，用户答案、正确答案、解析、笔记，错误答案标红
2. 错题导出（Word/PDF）：错题集，历史作答记录、答错次数、正确答案、解析、笔记
3. 题库导出（Word/PDF/Excel）：Word/PDF同错题格式，Excel同导入模板格式
所有导出均携带笔记部分
"""
import os
import sys
import json
from datetime import datetime

# 题型名称映射
TYPE_NAMES = {
    "single": "单选题",
    "multiple": "多选题",
    "judge": "判断题",
    "case": "案例分析题",
}


def _format_options(options):
    """格式化选项列表为 A. xxx 形式"""
    lines = []
    for i, opt in enumerate(options):
        letter = chr(65 + i)
        lines.append(f"{letter}. {opt}")
    return "\n".join(lines)


def _format_user_answer(user_answer, correct_answer, is_correct):
    """格式化用户答案，错误的标红标记"""
    if is_correct:
        return f"你的答案：{user_answer}"
    else:
        return f"你的答案：{user_answer}（错误）"


# ============ Word 导出 ============

def export_exam_record_word(exam_info, questions_data, output_path):
    """导出考试记录为Word
    exam_info: {exam_time, total, correct, wrong, score, duration}
    questions_data: [{index, q_type, question_text, options, user_answer, correct_answer, analysis, note_text, is_correct}, ...]
    """
    from docx import Document
    from docx.shared import Pt, RGBColor, Inches
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    doc = Document()

    # 设置默认字体
    style = doc.styles['Normal']
    font = style.font
    font.name = '宋体'
    font.size = Pt(11)

    # 标题
    title = doc.add_heading('考试记录', level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # 考试信息
    info = doc.add_paragraph()
    mins, secs = divmod(exam_info.get('duration', 0), 60)
    info.add_run(f"考试时间：{exam_info.get('exam_time', '')}\n")
    info.add_run(f"总题数：{exam_info.get('total', 0)} ｜ ")
    info.add_run(f"答对：{exam_info.get('correct', 0)} ｜ ")
    info.add_run(f"答错：{exam_info.get('wrong', 0)} ｜ ")
    info.add_run(f"得分：{exam_info.get('score', 0):.1f} ｜ ")
    info.add_run(f"用时：{mins:02d}:{secs:02d}")

    doc.add_paragraph('=' * 50)

    # 逐题
    for q in questions_data:
        idx = q.get('index', 0) + 1
        q_type = TYPE_NAMES.get(q.get('q_type', ''), q.get('q_type', ''))

        # 题号和题型
        p = doc.add_paragraph()
        run = p.add_run(f"第{idx}题【{q_type}】")
        run.bold = True
        run.font.size = Pt(12)

        # 题干
        doc.add_paragraph(f"题干：{q.get('question_text', '')}")

        # 选项
        options = q.get('options', [])
        if options:
            doc.add_paragraph(_format_options(options))

        doc.add_paragraph()

        # 用户答案（错误标红）
        p = doc.add_paragraph()
        if q.get('is_correct', False):
            p.add_run(f"你的答案：{q.get('user_answer', '')}")
        else:
            run = p.add_run(f"你的答案：{q.get('user_answer', '')}（错误）")
            run.font.color.rgb = RGBColor(200, 0, 0)

        # 正确答案
        doc.add_paragraph(f"正确答案：{q.get('correct_answer', '')}")

        # 解析
        if q.get('analysis'):
            doc.add_paragraph(f"解析：{q.get('analysis', '')}")

        # 笔记
        if q.get('note_text'):
            p = doc.add_paragraph()
            run = p.add_run(f"我的笔记：{q.get('note_text', '')}")
            run.font.color.rgb = RGBColor(0, 100, 0)

        doc.add_paragraph('-' * 50)

    doc.save(output_path)
    return output_path


def export_wrong_questions_word(wrong_data, output_path):
    """导出错题集为Word
    wrong_data: [{index, q_type, question_text, options, wrong_count, history_answers, correct_answer, analysis, note_text}, ...]
    history_answers: [{user_answer, created_at}, ...]
    """
    from docx import Document
    from docx.shared import Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    doc = Document()
    style = doc.styles['Normal']
    font = style.font
    font.name = '宋体'
    font.size = Pt(11)

    title = doc.add_heading('错题集导出', level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    info = doc.add_paragraph()
    info.add_run(f"导出时间：{datetime.now().strftime('%Y-%m-%d %H:%M:%S')}\n")
    info.add_run(f"错题总数：{len(wrong_data)}")

    doc.add_paragraph('=' * 50)

    for q in wrong_data:
        idx = q.get('index', 0) + 1
        q_type = TYPE_NAMES.get(q.get('q_type', ''), q.get('q_type', ''))
        wrong_count = q.get('wrong_count', 0)

        p = doc.add_paragraph()
        run = p.add_run(f"第{idx}题｜{q_type}｜累计答错：{wrong_count}次")
        run.bold = True
        run.font.size = Pt(12)

        doc.add_paragraph(f"题干：{q.get('question_text', '')}")

        options = q.get('options', [])
        if options:
            doc.add_paragraph(_format_options(options))

        doc.add_paragraph()

        # 历史作答记录
        doc.add_paragraph("历史作答记录：")
        history = q.get('history_answers', [])
        for i, h in enumerate(history):
            p = doc.add_paragraph()
            p.add_run(f"  第{i+1}次作答：{h.get('user_answer', '')}（错误）")
            p.runs[0].font.color.rgb = RGBColor(200, 0, 0)

        doc.add_paragraph()

        doc.add_paragraph(f"正确答案：{q.get('correct_answer', '')}")

        if q.get('analysis'):
            doc.add_paragraph(f"解析：{q.get('analysis', '')}")

        if q.get('note_text'):
            p = doc.add_paragraph()
            run = p.add_run(f"我的笔记：{q.get('note_text', '')}")
            run.font.color.rgb = RGBColor(0, 100, 0)

        doc.add_paragraph('-' * 50)

    doc.save(output_path)
    return output_path


def export_question_bank_word(questions, output_path):
    """导出题库为Word（同错题格式，去掉历史作答）"""
    from docx import Document
    from docx.shared import Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    doc = Document()
    style = doc.styles['Normal']
    font = style.font
    font.name = '宋体'
    font.size = Pt(11)

    title = doc.add_heading('题库导出', level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    info = doc.add_paragraph()
    info.add_run(f"导出时间：{datetime.now().strftime('%Y-%m-%d %H:%M:%S')}\n")
    info.add_run(f"题目总数：{len(questions)}")

    doc.add_paragraph('=' * 50)

    for i, q in enumerate(questions):
        idx = i + 1
        q_type = TYPE_NAMES.get(q.get('q_type', ''), q.get('q_type', ''))

        p = doc.add_paragraph()
        run = p.add_run(f"第{idx}题｜{q_type}")
        run.bold = True
        run.font.size = Pt(12)

        doc.add_paragraph(f"题干：{q.get('question_text', '')}")

        options = q.get('options', [])
        if options:
            doc.add_paragraph(_format_options(options))

        doc.add_paragraph()
        doc.add_paragraph(f"正确答案：{q.get('answer', '')}")

        if q.get('analysis'):
            doc.add_paragraph(f"解析：{q.get('analysis', '')}")

        if q.get('note_text'):
            p = doc.add_paragraph()
            run = p.add_run(f"我的笔记：{q.get('note_text', '')}")
            run.font.color.rgb = RGBColor(0, 100, 0)

        doc.add_paragraph('-' * 50)

    doc.save(output_path)
    return output_path


# ============ PDF 导出 ============

def _get_chinese_font():
    """获取支持中文的字体，优先使用系统字体"""
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    font_name = "ChineseFont"
    # 尝试注册系统中文字体
    if sys.platform == "win32":
        windir = os.environ.get("WINDIR", "C:\\Windows")
        font_paths = [
            os.path.join(windir, "Fonts", "simsun.ttc"),
            os.path.join(windir, "Fonts", "msyh.ttc"),
            os.path.join(windir, "Fonts", "simhei.ttf"),
            os.path.join(windir, "Fonts", "simkai.ttf"),
        ]
    else:
        font_paths = [
            "/usr/share/fonts/truetype/wqy/wqy-zenhei.ttc",
            "/usr/share/fonts/truetype/wqy/wqy-microhei.ttc",
            "/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc",
            "/System/Library/Fonts/PingFang.ttc",
            "/Library/Fonts/Arial Unicode.ttf",
            # 安卓系统字体
            "/system/fonts/NotoSansCJK-Regular.ttc",
            "/system/fonts/NotoSansSC-Regular.otf",
            "/system/fonts/DroidSansFallback.ttf",
        ]

    for fp in font_paths:
        if os.path.exists(fp):
            try:
                pdfmetrics.registerFont(TTFont(font_name, fp))
                return font_name
            except Exception:
                continue
    return "Helvetica"


def export_exam_record_pdf(exam_info, questions_data, output_path):
    """导出考试记录为PDF"""
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import cm
    from reportlab.lib.colors import red, green, black
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer

    font_name = _get_chinese_font()
    doc = SimpleDocTemplate(output_path, pagesize=A4,
                            leftMargin=2*cm, rightMargin=2*cm,
                            topMargin=2*cm, bottomMargin=2*cm)

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle('Title', parent=styles['Title'], fontName=font_name, fontSize=18, alignment=1)
    normal_style = ParagraphStyle('Normal', parent=styles['Normal'], fontName=font_name, fontSize=10, leading=14)
    bold_style = ParagraphStyle('Bold', parent=normal_style, fontName=font_name, fontSize=11, textColor=black)
    red_style = ParagraphStyle('Red', parent=normal_style, textColor=red)
    note_style = ParagraphStyle('Note', parent=normal_style, textColor=green)

    story = []
    story.append(Paragraph("考试记录", title_style))
    story.append(Spacer(1, 0.3*cm))

    mins, secs = divmod(exam_info.get('duration', 0), 60)
    info_text = (f"考试时间：{exam_info.get('exam_time', '')}<br/>"
                 f"总题数：{exam_info.get('total', 0)} ｜ 答对：{exam_info.get('correct', 0)} ｜ "
                 f"答错：{exam_info.get('wrong', 0)} ｜ 得分：{exam_info.get('score', 0):.1f} ｜ "
                 f"用时：{mins:02d}:{secs:02d}")
    story.append(Paragraph(info_text, normal_style))
    story.append(Spacer(1, 0.3*cm))
    story.append(Paragraph("=" * 60, normal_style))
    story.append(Spacer(1, 0.3*cm))

    for q in questions_data:
        idx = q.get('index', 0) + 1
        q_type = TYPE_NAMES.get(q.get('q_type', ''), q.get('q_type', ''))

        story.append(Paragraph(f"<b>第{idx}题【{q_type}】</b>", bold_style))
        story.append(Paragraph(f"题干：{q.get('question_text', '')}", normal_style))

        options = q.get('options', [])
        if options:
            story.append(Paragraph(_format_options(options).replace('\n', '<br/>'), normal_style))

        story.append(Spacer(1, 0.2*cm))

        if q.get('is_correct', False):
            story.append(Paragraph(f"你的答案：{q.get('user_answer', '')}", normal_style))
        else:
            story.append(Paragraph(f"你的答案：{q.get('user_answer', '')}（错误）", red_style))

        story.append(Paragraph(f"正确答案：{q.get('correct_answer', '')}", normal_style))

        if q.get('analysis'):
            story.append(Paragraph(f"解析：{q.get('analysis', '')}", normal_style))

        if q.get('note_text'):
            story.append(Paragraph(f"我的笔记：{q.get('note_text', '')}", note_style))

        story.append(Spacer(1, 0.2*cm))
        story.append(Paragraph("-" * 60, normal_style))
        story.append(Spacer(1, 0.3*cm))

    doc.build(story)
    return output_path


def export_wrong_questions_pdf(wrong_data, output_path):
    """导出错题集为PDF"""
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import cm
    from reportlab.lib.colors import red, green, black
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer

    font_name = _get_chinese_font()
    doc = SimpleDocTemplate(output_path, pagesize=A4,
                            leftMargin=2*cm, rightMargin=2*cm,
                            topMargin=2*cm, bottomMargin=2*cm)

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle('Title', parent=styles['Title'], fontName=font_name, fontSize=18, alignment=1)
    normal_style = ParagraphStyle('Normal', parent=styles['Normal'], fontName=font_name, fontSize=10, leading=14)
    bold_style = ParagraphStyle('Bold', parent=normal_style, fontName=font_name, fontSize=11)
    red_style = ParagraphStyle('Red', parent=normal_style, textColor=red)
    note_style = ParagraphStyle('Note', parent=normal_style, textColor=green)

    story = []
    story.append(Paragraph("错题集导出", title_style))
    story.append(Spacer(1, 0.3*cm))
    story.append(Paragraph(f"导出时间：{datetime.now().strftime('%Y-%m-%d %H:%M:%S')}<br/>错题总数：{len(wrong_data)}", normal_style))
    story.append(Spacer(1, 0.3*cm))
    story.append(Paragraph("=" * 60, normal_style))
    story.append(Spacer(1, 0.3*cm))

    for q in wrong_data:
        idx = q.get('index', 0) + 1
        q_type = TYPE_NAMES.get(q.get('q_type', ''), q.get('q_type', ''))
        wrong_count = q.get('wrong_count', 0)

        story.append(Paragraph(f"<b>第{idx}题｜{q_type}｜累计答错：{wrong_count}次</b>", bold_style))
        story.append(Paragraph(f"题干：{q.get('question_text', '')}", normal_style))

        options = q.get('options', [])
        if options:
            story.append(Paragraph(_format_options(options).replace('\n', '<br/>'), normal_style))

        story.append(Spacer(1, 0.2*cm))
        story.append(Paragraph("历史作答记录：", bold_style))

        history = q.get('history_answers', [])
        for i, h in enumerate(history):
            story.append(Paragraph(f"&nbsp;&nbsp;第{i+1}次作答：{h.get('user_answer', '')}（错误）", red_style))

        story.append(Spacer(1, 0.2*cm))
        story.append(Paragraph(f"正确答案：{q.get('correct_answer', '')}", normal_style))

        if q.get('analysis'):
            story.append(Paragraph(f"解析：{q.get('analysis', '')}", normal_style))

        if q.get('note_text'):
            story.append(Paragraph(f"我的笔记：{q.get('note_text', '')}", note_style))

        story.append(Spacer(1, 0.2*cm))
        story.append(Paragraph("-" * 60, normal_style))
        story.append(Spacer(1, 0.3*cm))

    doc.build(story)
    return output_path


def export_question_bank_pdf(questions, output_path):
    """导出题库为PDF"""
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import cm
    from reportlab.lib.colors import green, black
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer

    font_name = _get_chinese_font()
    doc = SimpleDocTemplate(output_path, pagesize=A4,
                            leftMargin=2*cm, rightMargin=2*cm,
                            topMargin=2*cm, bottomMargin=2*cm)

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle('Title', parent=styles['Title'], fontName=font_name, fontSize=18, alignment=1)
    normal_style = ParagraphStyle('Normal', parent=styles['Normal'], fontName=font_name, fontSize=10, leading=14)
    bold_style = ParagraphStyle('Bold', parent=normal_style, fontName=font_name, fontSize=11)
    note_style = ParagraphStyle('Note', parent=normal_style, textColor=green)

    story = []
    story.append(Paragraph("题库导出", title_style))
    story.append(Spacer(1, 0.3*cm))
    story.append(Paragraph(f"导出时间：{datetime.now().strftime('%Y-%m-%d %H:%M:%S')}<br/>题目总数：{len(questions)}", normal_style))
    story.append(Spacer(1, 0.3*cm))
    story.append(Paragraph("=" * 60, normal_style))
    story.append(Spacer(1, 0.3*cm))

    for i, q in enumerate(questions):
        idx = i + 1
        q_type = TYPE_NAMES.get(q.get('q_type', ''), q.get('q_type', ''))

        story.append(Paragraph(f"<b>第{idx}题｜{q_type}</b>", bold_style))
        story.append(Paragraph(f"题干：{q.get('question_text', '')}", normal_style))

        options = q.get('options', [])
        if options:
            story.append(Paragraph(_format_options(options).replace('\n', '<br/>'), normal_style))

        story.append(Spacer(1, 0.2*cm))
        story.append(Paragraph(f"正确答案：{q.get('answer', '')}", normal_style))

        if q.get('analysis'):
            story.append(Paragraph(f"解析：{q.get('analysis', '')}", normal_style))

        if q.get('note_text'):
            story.append(Paragraph(f"我的笔记：{q.get('note_text', '')}", note_style))

        story.append(Spacer(1, 0.2*cm))
        story.append(Paragraph("-" * 60, normal_style))
        story.append(Spacer(1, 0.3*cm))

    doc.build(story)
    return output_path


# ============ Excel 导出 ============

def export_question_bank_excel(questions, output_path):
    """导出题库为Excel，格式同导入模板，额外增加作答次数、错误次数、笔记列"""
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "题库"

    # 表头（同导入模板 + 额外列）
    headers = ["题型", "题干", "A选项", "B选项", "C选项", "D选项", "E选项", "F选项",
               "标准答案", "解析", "科目", "作答次数", "错误次数", "我的笔记"]
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF")

    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    # 题型映射
    type_map = {"single": "单选", "multiple": "多选", "judge": "判断", "case": "案例分析"}

    for row_idx, q in enumerate(questions, 2):
        ws.cell(row=row_idx, column=1, value=type_map.get(q.get("q_type", ""), q.get("q_type", "")))
        ws.cell(row=row_idx, column=2, value=q.get("question_text", ""))

        options = q.get("options", [])
        for i in range(6):
            if i < len(options):
                ws.cell(row=row_idx, column=3 + i, value=options[i])

        ws.cell(row=row_idx, column=9, value=q.get("answer", ""))
        ws.cell(row=row_idx, column=10, value=q.get("analysis", ""))
        # 科目名称需要查询，这里留空或用subject_id
        ws.cell(row=row_idx, column=11, value="")
        ws.cell(row=row_idx, column=12, value=q.get("try_count", 0))
        ws.cell(row=row_idx, column=13, value=q.get("wrong_count", 0))
        ws.cell(row=row_idx, column=14, value=q.get("note_text", ""))

    # 调整列宽
    col_widths = [10, 40, 15, 15, 15, 15, 15, 15, 12, 30, 12, 10, 10, 30]
    for i, w in enumerate(col_widths, 1):
        ws.column_dimensions[chr(64 + i)].width = w

    wb.save(output_path)
    return output_path
